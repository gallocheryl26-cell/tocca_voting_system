from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path

doc = Document(r"c:\Users\Rysha\Downloads\Capstone proposal 1 ROMO GROUP (revised) CLEAN.docx")
print("sections", len(doc.sections))
for i, s in enumerate(doc.sections):
    print(f"sec{i} L={s.left_margin.inches:.2f} R={s.right_margin.inches:.2f} T={s.top_margin.inches:.2f} B={s.bottom_margin.inches:.2f} first={s.different_first_page_header_footer} start={s.start_type}")
    print("  footer:", [p.text for p in s.footer.paragraphs])

align_map = {0:"LEFT",1:"CENTER",2:"RIGHT",3:"JUSTIFY",None:"None"}
print("\n--- key paras ---")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    if i > 80 and t not in (
        "Introduction","Project Context","Purpose and Description","Objectives","General Objective:",
        "Specific Objectives:","Scope and Limitations","Limitations",
        "Review of Related Literature/Studies/Systems","METHODOLOGY","Technical Background",
        "Resources","Requirements Analysis","Requirements Documentation","Design Model",
        "Design of Software, System, Product, and/or Processes","Development",
        "RESULTS AND DISCUSSION","Testing","Description of Prototype","Implementation Plan",
        "CONCLUSION","References","APPENDICES","APPENDIX A. RESOURCE PERSONS"
    ) and not t.startswith("APPENDIX"):
        continue
    if i <= 45 or t in (
        "Introduction","Project Context","Purpose and Description","Objectives","METHODOLOGY",
        "Technical Background","Resources","Requirements Analysis","Design Model",
        "RESULTS AND DISCUSSION","CONCLUSION","References","APPENDICES"
    ) or t.startswith("APPENDIX") or t.startswith("ENDORSEMENT") or t.startswith("APPROVAL") or t in ("Acknowledgements","Abstract","Table of Contents"):
        al = align_map.get(p.alignment, str(p.alignment))
        print(f"{i:03d} {p.style.name!r:22} al={al:7} | {t[:70]}")
