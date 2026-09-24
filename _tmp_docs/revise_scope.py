# -*- coding: utf-8 -*-
"""Replace Scope and Limitations with per-user, per-module STI hardbound format."""
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

SRC = Path(r"C:\Users\Rysha\OneDrive - ormoc.sti.ph\Desktop\PRACTICOM (4) from PDF.docx")
OUT = Path(r"C:\Users\Rysha\OneDrive - ormoc.sti.ph\Desktop\PRACTICOM (4) Scope Revised.docx")

# kind: heading | group | module | body | limit_title | bullet
CONTENT = [
    ("body",
     "This capstone project aims to enhance On-the-Job Training at STI College Ormoc by developing a more secure, efficient, and user-friendly Digital Practicum and Internship Deployment and Monitoring System. The study is generally useful in strengthening record accuracy, improving coordination, and simplifying the overall management of practicum deployment. The main beneficiaries of the system include the OJT coordinators and super administrators of STI College Ormoc as organizers, students in the BSIT, BSTM, and BSHM programs as interns, and verified industry partners as host companies."),
    ("body",
     "The project was conducted during Academic Year 2025–2026, with development from January 2026 through final submission on May 15, 2026. The system supports three academic programs—BSIT, BSTM, and BSHM—with course-specific hour requirements of 486 hours for BSIT and 600 hours for BSTM and BSHM. Testing and evaluation are limited to the intended users and workflows of STI College Ormoc’s OJT program. The backend is built using PHP and MySQL with a REST API connecting the mobile client, deployed on XAMPP for local development and on practicom.online for production use. External integrations within scope include Microsoft Office 365 OAuth for student authentication, Face++ for biometric attendance, Semaphore for SMS verification, and Maileroo for partner onboarding email. The system also includes the modules:"),
    ("group", "Super Administrator (Web Panel):"),
    ("module", "Login Module – contains signing in to the administrative web panel and restricting access to authorized super administrators."),
    ("module", "Dashboard – contains institution-wide summaries of applications, deployments, attendance, documents, and program activity."),
    ("module", "User Management – contains the adding, updating, and deactivating of coordinator and administrator accounts."),
    ("module", "Partner Management – contains the review, approval, and rejection of industry partner registrations and verification documents across programs."),
    ("module", "Internship Posting Oversight – contains viewing and overseeing internship postings published for BSIT, BSTM, and BSHM."),
    ("module", "Application and Deployment Oversight – contains viewing student applications and deployed interns across programs."),
    ("module", "Reports and Analytics – contains dashboards and exportable summaries of applications, deployments, attendance, and rendered hours."),
    ("group", "OJT Coordinator (Web Panel):"),
    ("module", "Login Module – contains signing in to the administrative web panel for the coordinator’s assigned program."),
    ("module", "Dashboard – contains program-level summaries of applications, deployments, attendance, and documents."),
    ("module", "Partner Management – contains the review, approval, and rejection of industry partners relevant to the assigned course."),
    ("module", "Internship Posting Management – contains adding, updating, and overseeing internship postings for the assigned program."),
    ("module", "Application Processing – contains moving student applications through pending, under review, for interview, for deployment, and deployed."),
    ("module", "Document Management – contains tracking Memoranda of Agreement, endorsement letters, and student-submitted files required for deployment."),
    ("module", "Attendance and Hours Monitoring – contains reviewing attendance logs and comparing accumulated hours with program hour requirements."),
    ("module", "Announcements and Calendar – contains posting announcements and calendar events for students and partners."),
    ("module", "Reports – contains program reports on applications, deployments, attendance, and hours."),
    ("group", "OJT Coordinator (Mobile Application):"),
    ("module", "Coordinator Login – contains signing in on mobile for field monitoring."),
    ("module", "Attendance Monitor – contains viewing deployed students and inspecting attendance history while away from the office."),
    ("module", "Student Progress View – contains selected student reports needed for on-site coordination."),
    ("group", "Student (Mobile Application):"),
    ("module", "Office 365 Login – contains signing in using the institutional Microsoft Office 365 account issued by STI College Ormoc."),
    ("module", "Profile Management – contains completing the student profile and uploading required documents."),
    ("module", "Internship Browse and Details – contains viewing course-matched internship postings and posting details."),
    ("module", "Application Module – contains submitting an application with résumé and supporting files and viewing application status."),
    ("module", "Attendance Module – contains clock-in and clock-out, optional facial recognition through Face++, and viewing attendance logs and hour summaries."),
    ("module", "Excuse Submission – contains submitting excuse requests for missed or incomplete duty hours."),
    ("module", "Notifications – contains receiving in-app notices on application status, announcements, and related updates."),
    ("group", "Industry Partner (Web Portal):"),
    ("module", "Login and Password Change – contains signing in to the company portal and completing the required initial password change."),
    ("module", "Verification Documents – contains uploading company verification files for school approval."),
    ("module", "Internship Postings – contains adding and managing placement opportunities after the company is approved."),
    ("module", "Applicant Review – contains viewing and reviewing student applications."),
    ("module", "Intern Management – contains viewing deployed interns assigned to the company."),
    ("module", "Time Log Validation – contains reviewing and validating intern attendance logs."),
    ("module", "Announcements and Calendar – contains viewing or posting announcements and calendar items as implemented."),
    ("group", "Public Users (Landing Page):"),
    ("module", "Landing Page – contains browsing active internship postings and featured industry partners prior to login."),
    ("limit_title", "Limitations"),
    ("bullet",
     "Platform coverage — The system does not include a dedicated mobile application for industry partners; host companies access the platform exclusively through the web portal, which may limit convenience for partners who prefer mobile-based supervision."),
    ("bullet",
     "Institutional dependency — Student authentication relies on Microsoft Office 365 accounts issued by STI College Ormoc. Users without valid institutional credentials cannot access the mobile application, restricting use to enrolled students and authorized personnel."),
    ("bullet",
     "Internet connectivity requirement — PRACTICOM requires a stable internet connection for all core functions, including attendance logging, face verification, and API communication. Offline mode is not supported."),
    ("bullet",
     "Third-party service dependency — Biometric attendance verification depends on the Face++ API, and partner onboarding relies on Semaphore and Maileroo. Service outages, API changes, or connectivity issues with these providers may affect system functionality."),
    ("bullet",
     "Password hashing algorithm — The current implementation stores web-based admin and company account passwords using MD5, a hashing algorithm that is now considered cryptographically broken and insecure for password storage: it is fast to compute, has no built-in salting, and is vulnerable to collision attacks and rainbow-table lookups. This is a known security limitation of the present implementation and not a recommended practice. Mobile accounts are not affected, as they use token-based (Bearer) authentication rather than stored password hashes. Migrating web account password storage to a salted, adaptive hashing algorithm such as bcrypt or Argon2, together with HTTPS enforcement and CSRF protection, is identified as a required step before any production or public deployment of the system beyond the scope of this study."),
]


def has_sectpr(p):
    pPr = p._p.find(qn("w:pPr"))
    if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
        return True
    return p._p.find(qn("w:sectPr")) is not None


def clear_runs(p):
    for child in list(p._p):
        if child.tag != qn("w:pPr"):
            p._p.remove(child)


def add_run(p, text, *, bold=False, size_pt=12):
    r = p.add_run(text)
    r.bold = bold
    r.italic = False
    r.font.name = "Times New Roman"
    r.font.size = __import__("docx.shared", fromlist=["Pt"]).Pt(size_pt)
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = __import__("docx.oxml", fromlist=["OxmlElement"]).OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:cs"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    return r


def clone_p_after(anchor_p, template_p):
    new_el = deepcopy(template_p._p)
    # strip sectPr if accidentally copied
    pPr = new_el.find(qn("w:pPr"))
    if pPr is not None:
        sect = pPr.find(qn("w:sectPr"))
        if sect is not None:
            pPr.remove(sect)
    for child in list(new_el):
        if child.tag != qn("w:pPr"):
            new_el.remove(child)
    anchor_p._p.addnext(new_el)
    return Paragraph(new_el, anchor_p._parent)


def main():
    from shutil import copy2
    copy2(SRC, OUT)
    doc = Document(str(OUT))

    def find_span(document):
        s = e = None
        for i, p in enumerate(document.paragraphs):
            t = p.text.strip()
            nxt = document.paragraphs[i + 1].text.strip() if i + 1 < len(document.paragraphs) else ""
            if "Scope and Limitations" in t and (
                nxt.startswith("This study covers") or nxt.startswith("This capstone project")
            ):
                s = i
            if s is not None and t.startswith("Review of Related Literature") and "\t" not in t:
                if nxt.startswith("Foreign") or nxt.startswith("Several studies") or nxt.startswith("This chapter"):
                    e = i
                    break
        return s, e

    extra_remove = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("The Header remove it") or t.startswith("Scope is per user"):
            extra_remove.append(p)
    for p in extra_remove:
        parent = p._p.getparent()
        if parent is not None:
            parent.remove(p._p)

    start, end = find_span(doc)
    if start is None or end is None:
        raise SystemExit(f"markers not found start={start} end={end}")

    heading = doc.paragraphs[start]
    clear_runs(heading)
    add_run(heading, "Scope and Limitations", bold=True)
    template = None
    for i in range(start + 1, end):
        p = doc.paragraphs[i]
        if p.text.strip() and not has_sectpr(p):
            template = p
            break
    if template is None:
        template = heading

    # delete old body paras between heading and RRL, keep sectPr shells
    to_remove = []
    for i in range(start + 1, end):
        p = doc.paragraphs[i]
        if has_sectpr(p):
            clear_runs(p)
        else:
            to_remove.append(p)
    for p in to_remove:
        parent = p._p.getparent()
        if parent is not None:
            parent.remove(p._p)

    # insert new content after heading
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    anchor = heading
    for kind, text in CONTENT:
        np = clone_p_after(anchor, template)
        if kind in ("body", "module", "bullet"):
            np.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        else:
            np.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if kind == "group":
            add_run(np, text, bold=True)
        elif kind == "limit_title":
            add_run(np, text, bold=True)
        elif kind == "module":
            name, rest = text.split(" – ", 1)
            add_run(np, name + " – ", bold=True)
            add_run(np, rest, bold=False)
        elif kind == "bullet":
            label, rest = text.split(" — ", 1)
            add_run(np, "• " + label + " — ", bold=True)
            add_run(np, rest, bold=False)
        else:
            add_run(np, text, bold=False)
        anchor = np

    doc.save(str(OUT))
    print("Wrote", OUT)

    # verify
    doc2 = Document(str(OUT))
    printing = False
    n = 0
    for p in doc2.paragraphs:
        t = p.text.strip()
        if t == "Scope and Limitations":
            printing = True
        if printing:
            print(t[:110] if t else "")
            n += 1
        if printing and t.startswith("Review of Related"):
            break
    print("lines printed", n)


if __name__ == "__main__":
    main()
