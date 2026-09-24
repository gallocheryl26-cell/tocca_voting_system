# -*- coding: utf-8 -*-
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from docx import Document
from docx.oxml.ns import qn

REF = Path(r"c:\Users\Rysha\OneDrive - ormoc.sti.ph\Documents\TOCCA Capstone Final Document Revision.docx")
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

with ZipFile(REF) as z:
    styles = ET.fromstring(z.read("word/styles.xml"))
    print("=== ALL STYLE NAMES ===")
    for st in styles.findall(f"{W}style"):
        name_el = st.find(f"{W}name")
        name = name_el.get(f"{W}val") if name_el is not None else "?"
        sid = st.get(f"{W}styleId")
        typ = st.get(f"{W}type")
        if typ == "paragraph":
            print(f"  {sid!r} name={name!r}")

    print("\n=== Heading1 / Heading2 / Heading3 full XML (truncated) ===")
    for st in styles.findall(f"{W}style"):
        sid = st.get(f"{W}styleId")
        if sid in ("Heading1", "Heading2", "Heading3", "Title", "Normal"):
            xml = ET.tostring(st, encoding="unicode")
            print(f"\n----- {sid} len={len(xml)} -----")
            print(xml[:2500])

    print("\n=== FOOTERS ===")
    for name in ["word/footer1.xml", "word/footer2.xml", "word/footer3.xml", "word/footer4.xml"]:
        print("\n", name)
        print(z.read(name).decode("utf-8", errors="ignore")[:2500])

doc = Document(str(REF))
print("\n=== TABLES (first 3) ===")
for i, t in enumerate(doc.tables[:3]):
    print(f"\nTABLE {i} {len(t.rows)}x{len(t.columns)}")
    for ri, row in enumerate(t.rows[:25]):
        cells = [c.text.replace("\n", " | ")[:80] for c in row.cells]
        print(f"  r{ri}: {cells}")

print("\n=== TITLE PAGE PARAS with spacing ===")
for i, p in enumerate(doc.paragraphs[:25]):
    pf = p.paragraph_format
    print(f"{i:02d} style={p.style.name!r} sp_before={pf.space_before} after={pf.space_after} ls={pf.line_spacing} ind={pf.first_line_indent} | {p.text[:70]!r}")
